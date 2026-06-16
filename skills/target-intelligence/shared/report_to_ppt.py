#!/usr/bin/env python3
"""
Report-to-PPTX: Generate a professional presentation from target analysis data.

Inspired by nature-paper2ppt's presentation arc:
- Target Overview (slide 1)
- Expression Profile (slide 2)
- Survival Analysis (slide 3)
- Immune Landscape (slide 4)
- Competitive Intelligence (slide 5)
- Development Timeline (slide 6)
- Key Findings & Recommendations (slide 7)

Usage:
    python report_to_ppt.py --gene KRAS --output-dir ./output/KRAS
"""

import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    HAS_PPTX = True

    # ── Color Palette (Nature-style, restrained) ──────────────────────────────────
    COLORS = {
        'bg': RGBColor(0x0F, 0x17, 0x2A),
        'bg2': RGBColor(0x1E, 0x29, 0x3B),
        'accent': RGBColor(0x3B, 0x82, 0xF6),
        'accent2': RGBColor(0x60, 0xA5, 0xFA),
        'success': RGBColor(0x22, 0xC5, 0x5E),
        'warn': RGBColor(0xF5, 0x9E, 0x0B),
        'danger': RGBColor(0xEF, 0x44, 0x44),
        'fg': RGBColor(0xE2, 0xE8, 0xF0),
        'fg2': RGBColor(0x94, 0xA3, 0xB8),
        'fg3': RGBColor(0x64, 0x74, 0x8B),
        'white': RGBColor(0xFF, 0xFF, 0xFF),
        'dark': RGBColor(0x0F, 0x17, 0x2A),
    }
except ImportError:
    HAS_PPTX = False
    COLORS = {}


def load_data(output_dir: Path, gene: str) -> dict:
    """Load all analysis data from output directory."""
    data = {}

    files = {
        'gtex': f'gtex_{gene}_expression.json',
        'survival': f'survival_*_results.json',
        'immune': f'immune_*_correlation.json',
        'competitive': f'competitive_{gene}_intelligence.json',
        'patent': f'patent_{gene}_landscape.json',
        'protein': f'protein_{gene}_validation.json',
        'milestones': f'milestones_{gene}.json',
    }

    for key, pattern in files.items():
        if '*' in pattern:
            matches = list(output_dir.glob(pattern))
            if matches:
                try:
                    data[key] = json.loads(matches[0].read_text(encoding='utf-8'))
                except Exception:
                    pass
        else:
            f = output_dir / pattern
            if f.exists():
                try:
                    data[key] = json.loads(f.read_text(encoding='utf-8'))
                except Exception:
                    pass

    return data


def add_styled_slide(prs, title_text, subtitle_text=''):
    """Add a slide with consistent styling."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Dark background
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['bg']

    # Title
    left = Inches(0.8)
    top = Inches(0.4)
    width = Inches(8.4)
    height = Inches(0.8)
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(28)
    p.font.color.rgb = COLORS['accent2']
    p.font.bold = True

    # Subtitle
    if subtitle_text:
        top = Inches(1.2)
        height = Inches(0.5)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = subtitle_text
        p.font.size = Pt(14)
        p.font.color.rgb = COLORS['fg3']

    # Accent line
    line_left = Inches(0.8)
    line_top = Inches(1.15)
    line = slide.shapes.add_shape(1, line_left, line_top, Inches(8.4), Pt(3))  # Rectangle
    line.fill.solid()
    line.fill.fore_color.rgb = COLORS['accent']
    line.line.fill.background()

    return slide


def add_metric_card(slide, left, top, width, height, label, value, sub=''):
    """Add a metric card to a slide."""
    # Card background
    card = slide.shapes.add_shape(1, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = COLORS['bg2']
    card.line.color.rgb = RGBColor(0x33, 0x41, 0x55)
    card.line.width = Pt(1)

    # Label
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), width - Inches(0.3), Inches(0.3))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = label.upper()
    p.font.size = Pt(9)
    p.font.color.rgb = COLORS['fg3']

    # Value
    txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.35), width - Inches(0.3), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = str(value)
    p.font.size = Pt(22)
    p.font.color.rgb = COLORS['accent2']
    p.font.bold = True

    # Sub
    if sub:
        txBox = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.8), width - Inches(0.3), Inches(0.3))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = sub
        p.font.size = Pt(9)
        p.font.color.rgb = COLORS['fg3']


def add_table(slide, left, top, width, headers, rows, font_size=10):
    """Add a styled table to a slide."""
    n_rows = min(len(rows) + 1, 10)  # +1 for header, max 10
    n_cols = len(headers)
    row_height = Inches(0.35)

    table_shape = slide.shapes.add_table(n_rows, n_cols, left, top, width, row_height * n_rows)
    table = table_shape.table

    # Style header
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS['accent']
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(font_size)
            paragraph.font.color.rgb = COLORS['white']
            paragraph.font.bold = True

    # Style rows
    for r_idx, row in enumerate(rows[:n_rows - 1]):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS['bg2'] if r_idx % 2 == 0 else COLORS['bg']
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(font_size)
                paragraph.font.color.rgb = COLORS['fg2']


def slide_title_overview(prs, gene, data):
    """Slide 1: Target overview with key metrics."""
    slide = add_styled_slide(prs, f'{gene} Target Intelligence Report', 'Comprehensive Drug Target Analysis')

    # Key metrics row
    cards = []
    protein = data.get('protein', {}).get('protein_info', {})
    if protein:
        cards.append(('Target', protein.get('protein_name', gene), f'UniProt: {protein.get("accession", "—")}'))

    gtex = data.get('gtex', {})
    stats = gtex.get('summary_stats', {})
    if stats:
        cards.append(('Tissues', str(stats.get('tissues_above_1_tpm', '—')), f'Tau: {stats.get("tau_index", 0):.3f}'))

    surv = data.get('survival', {}).get('survival_stats', {})
    if surv:
        hr = surv.get('cox_regression', {}).get('hazard_ratio', '—')
        cards.append(('Hazard Ratio', f'{hr:.3f}' if isinstance(hr, (int, float)) else str(hr),
                       f'P = {surv.get("logrank_pvalue", 0):.2e}'))

    comp = data.get('competitive', {}).get('competition', {})
    if comp:
        cards.append(('Clinical Trials', str(comp.get('total_trials', '—')), 'Active trials'))

    pat = data.get('patent', {}).get('patent_trends', {})
    if pat:
        cards.append(('Patents', str(pat.get('total_patents', '—')), 'Related patents'))

    card_width = Inches(1.7)
    card_height = Inches(1.2)
    start_left = Inches(0.8)
    gap = Inches(0.15)

    for i, (label, value, sub) in enumerate(cards[:5]):
        left = start_left + (card_width + gap) * i
        add_metric_card(slide, left, Inches(1.8), card_width, card_height, label, value, sub)

    # Function description
    func = protein.get('function', '')
    if func:
        if len(func) > 250:
            func = func[:250] + '...'
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(3.3), Inches(8.4), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = func
        p.font.size = Pt(12)
        p.font.color.rgb = COLORS['fg2']
        p.line_spacing = Pt(18)


def slide_expression(prs, gene, data):
    """Slide 2: Expression profile."""
    gtex = data.get('gtex', {})
    expr = gtex.get('expression_by_tissue', {})
    if not expr:
        return

    slide = add_styled_slide(prs, 'Tissue Expression Profile', f'{gene} — GTEx RNA-seq')

    # Top 10 tissues table
    sorted_tissues = sorted(expr.items(), key=lambda x: x[1], reverse=True)[:10]
    headers = ['Rank', 'Tissue', 'TPM', 'Relative']
    max_tpm = sorted_tissues[0][1] if sorted_tissues else 1
    rows = []
    for i, (tissue, tpm) in enumerate(sorted_tissues):
        rel = '█' * int(tpm / max_tpm * 20)
        rows.append([str(i + 1), tissue, f'{tpm:.1f}', rel])

    add_table(slide, Inches(0.8), Inches(1.6), Inches(5), headers, rows)

    # Stats card
    stats = gtex.get('summary_stats', {})
    specificity = gtex.get('tissue_specificity', {})
    tau = stats.get('tau_index') or specificity.get('tau_index', 0)
    level = stats.get('specificity_level') or specificity.get('specificity_level', '')

    add_metric_card(slide, Inches(6.2), Inches(1.6), Inches(2.6), Inches(1.2),
                    'Tau Index', f'{tau:.3f}', level)
    add_metric_card(slide, Inches(6.2), Inches(3.0), Inches(2.6), Inches(1.2),
                    'Total Tissues', str(len(expr)), 'detected tissues')


def slide_survival(prs, gene, data):
    """Slide 3: Survival analysis."""
    surv_data = data.get('survival', {}).get('survival_stats', {})
    if not surv_data:
        return

    slide = add_styled_slide(prs, 'Prognostic Value', f'{gene} — TCGA Survival Analysis')

    cox = surv_data.get('cox_regression', {})

    # Metrics
    cards = [
        ('Hazard Ratio', f'{cox.get("hazard_ratio", 0):.3f}',
         f'95% CI: {cox.get("hr_ci_lower", 0):.2f} - {cox.get("hr_ci_upper", 0):.2f}'),
        ('Cox P-value', f'{cox.get("cox_pvalue", 0):.2e}', ''),
        ('Log-rank P', f'{surv_data.get("logrank_pvalue", 0):.2e}', ''),
        ('Median High', f'{surv_data.get("median_survival_high_months", 0):.1f} mo',
         f'{surv_data.get("events_high", 0)} events / {surv_data.get("n_high", 0)} pts'),
        ('Median Low', f'{surv_data.get("median_survival_low_months", 0):.1f} mo',
         f'{surv_data.get("events_low", 0)} events / {surv_data.get("n_low", 0)} pts'),
    ]

    for i, (label, value, sub) in enumerate(cards):
        left = Inches(0.8) + Inches(1.75) * i
        add_metric_card(slide, left, Inches(1.6), Inches(1.6), Inches(1.2), label, value, sub)

    # Interpretation
    hr = cox.get('hazard_ratio', 1)
    pval = cox.get('cox_pvalue', 1)
    if hr > 1 and pval < 0.05:
        interp = f'{gene} high expression is associated with significantly worse survival (HR={hr:.3f}, P={pval:.2e}). This suggests {gene} expression is a negative prognostic marker in this cancer type.'
        color = COLORS['danger']
    elif hr < 1 and pval < 0.05:
        interp = f'{gene} high expression is associated with better survival (HR={hr:.3f}, P={pval:.2e}). This suggests a protective role.'
        color = COLORS['success']
    else:
        interp = f'Survival association is not statistically significant (HR={hr:.3f}, P={pval:.2e}). Further validation needed.'
        color = COLORS['warn']

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(3.2), Inches(8.4), Inches(1.5))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = interp
    p.font.size = Pt(13)
    p.font.color.rgb = color
    p.line_spacing = Pt(20)


def slide_immune(prs, gene, data):
    """Slide 4: Immune landscape."""
    immune = data.get('immune', {}).get('immune_correlations', {})
    if not immune:
        return

    slide = add_styled_slide(prs, 'Immune Microenvironment', f'{gene} — Tumor Immune Correlation')

    headers = ['Cell Type', 'Correlation (r)', 'P-value', 'Adj. P', 'Significant']
    rows = []
    for cell, vals in sorted(immune.items(), key=lambda x: x[1].get('r', 0)):
        r = vals.get('r', 0)
        pval = vals.get('pvalue', 0)
        adj = vals.get('adj_pvalue', 0)
        sig = 'Yes' if vals.get('significant') else 'No'
        rows.append([cell, f'{r:.3f}', f'{pval:.2e}', f'{adj:.3f}', sig])

    add_table(slide, Inches(0.8), Inches(1.6), Inches(8.4), headers, rows)

    # Summary
    n_sig = sum(1 for v in immune.values() if v.get('significant'))
    n_total = len(immune)
    neg = [k for k, v in immune.items() if v.get('r', 0) < -0.1 and v.get('significant')]
    pos = [k for k, v in immune.items() if v.get('r', 0) > 0.1 and v.get('significant')]

    summary = f'{n_sig}/{n_total} cell types significantly correlated.'
    if neg:
        summary += f' Negative: {", ".join(neg)}.'
    if pos:
        summary += f' Positive: {", ".join(pos)}.'

    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(8.4), Inches(0.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = summary
    p.font.size = Pt(12)
    p.font.color.rgb = COLORS['fg2']


def slide_competitive(prs, gene, data):
    """Slide 5: Competitive intelligence."""
    comp = data.get('competitive', {}).get('competition', {})
    if not comp:
        return

    slide = add_styled_slide(prs, 'Competitive Landscape', f'{gene} — Clinical Pipeline & Drugs')

    # Approved drugs
    approved = comp.get('approved_drugs', [])
    if approved:
        headers = ['Drug', 'Mechanism', 'Approval']
        rows = [[d.get('name', ''), d.get('mechanism', ''), d.get('first_approval', '')] for d in approved[:4]]
        add_table(slide, Inches(0.8), Inches(1.6), Inches(5), headers, rows)

    # Trials by phase
    pd = comp.get('phase_distribution', {})
    if pd:
        headers = ['Phase', 'Count']
        rows = [[phase, str(count)] for phase, count in sorted(pd.items())]
        add_table(slide, Inches(6.2), Inches(1.6), Inches(3), headers, rows)

    # Key mechanisms
    mechs = comp.get('key_mechanisms', [])
    if mechs:
        mech_text = ', '.join([f"{m.get('mechanism', '')} ({m.get('count', 0)})" for m in mechs[:3]])
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(8.4), Inches(0.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f'Key Mechanisms: {mech_text}'
        p.font.size = Pt(12)
        p.font.color.rgb = COLORS['fg2']

    # Total counts
    pat = data.get('patent', {}).get('patent_trends', {})
    add_metric_card(slide, Inches(0.8), Inches(4.2), Inches(2.5), Inches(1.0),
                    'Total Trials', str(comp.get('total_trials', 0)), '')
    add_metric_card(slide, Inches(3.6), Inches(4.2), Inches(2.5), Inches(1.0),
                    'Approved Drugs', str(len(approved)), '')
    add_metric_card(slide, Inches(6.4), Inches(4.2), Inches(2.5), Inches(1.0),
                    'Total Patents', str(pat.get('total_patents', 0)), '')


def slide_milestones(prs, gene, data):
    """Slide 6: Development timeline."""
    ms = data.get('milestones', {})
    milestones = ms.get('milestones', [])
    if not milestones:
        return

    slide = add_styled_slide(prs, 'Development Timeline', f'{gene} — {ms.get("modality", "Small Molecule")}')

    headers = ['Phase', 'Start', 'End', 'Duration']
    rows = []
    for m in milestones:
        name = m.get('phase', m.get('name', ''))
        start = m.get('start_date', m.get('start', ''))
        end = m.get('end_date', m.get('end', ''))
        dur = f'{m.get("duration_months", 0)} mo'
        rows.append([name, start, end, dur])

    add_table(slide, Inches(0.8), Inches(1.6), Inches(8.4), headers, rows)

    total = ms.get('total_timeline_months', '')
    if isinstance(total, str):
        total = sum(m.get('duration_months', 0) for m in milestones)
    add_metric_card(slide, Inches(0.8), Inches(4.2), Inches(2.5), Inches(1.0),
                    'Total Timeline', f'{total} months', f'~{total/12:.1f} years')


def slide_recommendations(prs, gene, data):
    """Slide 7: Key findings and recommendations."""
    slide = add_styled_slide(prs, 'Key Findings & Recommendations', gene)

    findings = []

    # Expression findings
    gtex = data.get('gtex', {})
    stats = gtex.get('summary_stats', {})
    tau = stats.get('tau_index', 0)
    if tau > 0.5:
        findings.append(f'Tissue-specific expression (Tau={tau:.3f}) — potential for targeted therapy')

    # Survival findings
    surv = data.get('survival', {}).get('survival_stats', {})
    cox = surv.get('cox_regression', {})
    hr = cox.get('hazard_ratio', 1)
    pval = cox.get('cox_pvalue', 1)
    if pval < 0.05:
        direction = 'worse' if hr > 1 else 'better'
        findings.append(f'Significant survival association (HR={hr:.3f}, P={pval:.2e}) — {direction} prognosis')

    # Immune findings
    immune = data.get('immune', {}).get('immune_correlations', {})
    sig_neg = [k for k, v in immune.items() if v.get('r', 0) < -0.1 and v.get('significant')]
    if sig_neg:
        findings.append(f'Negative immune correlation ({", ".join(sig_neg)}) — potential immune evasion')

    # Competition findings
    comp = data.get('competitive', {}).get('competition', {})
    n_trials = comp.get('total_trials', 0)
    if n_trials > 0:
        findings.append(f'{n_trials} active clinical trials — active development landscape')

    # Render findings
    txBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(8.4), Inches(3.5))
    tf = txBox.text_frame
    tf.word_wrap = True

    for i, finding in enumerate(findings):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f'  {finding}'
        p.font.size = Pt(14)
        p.font.color.rgb = COLORS['fg']
        p.space_after = Pt(12)
        p.line_spacing = Pt(22)


def generate_pptx(gene: str, output_dir: Path) -> Path:
    """Generate the PPTX presentation."""
    if not HAS_PPTX:
        print('[PPT] python-pptx not installed. Run: pip install python-pptx')
        return None

    data = load_data(output_dir, gene)
    if not data:
        print(f'[PPT] No data found for {gene} in {output_dir}')
        return None

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9

    slide_title_overview(prs, gene, data)
    slide_expression(prs, gene, data)
    slide_survival(prs, gene, data)
    slide_immune(prs, gene, data)
    slide_competitive(prs, gene, data)
    slide_milestones(prs, gene, data)
    slide_recommendations(prs, gene, data)

    # Footer on all slides
    for slide in prs.slides:
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(8.4), Inches(0.3))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = f'Target Intelligence Report | {gene} | Auto-generated'
        p.font.size = Pt(8)
        p.font.color.rgb = COLORS['fg3']
        p.alignment = PP_ALIGN.RIGHT

    out_path = output_dir / f'{gene}_presentation.pptx'
    prs.save(str(out_path))
    print(f'[PPT] Saved: {out_path}')
    return out_path


def main():
    parser = argparse.ArgumentParser(description='Generate PPTX from target analysis data')
    parser.add_argument('--gene', required=True, help='Gene symbol')
    parser.add_argument('--output-dir', required=True, help='Output directory')
    args = parser.parse_args()

    generate_pptx(args.gene, Path(args.output_dir))


if __name__ == '__main__':
    main()
