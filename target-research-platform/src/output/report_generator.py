"""
报告生成器
生成 HTML / PDF / DOCX 研究报告
"""
import json
from pathlib import Path
from typing import Optional
from datetime import datetime

from jinja2 import Environment, FileSystemLoader, BaseLoader


# HTML 报告模板（内联，不依赖外部文件）
HTML_TEMPLATE = r"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{{ title }}</title>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body { font-family: -apple-system, 'Segoe UI', Roboto, 'SimHei', sans-serif; background: #f5f7fa; color: #2c3e50; line-height: 1.6; }
        .container { max-width: 1200px; margin: 0 auto; padding: 20px; }
        .header { background: linear-gradient(135deg, #1a5276, #2980b9); color: white; padding: 40px; border-radius: 12px; margin-bottom: 24px; }
        .header h1 { font-size: 28px; margin-bottom: 8px; }
        .header .meta { opacity: 0.85; font-size: 14px; }
        .grid { display: grid; grid-template-columns: 1fr 1fr; gap: 20px; margin-bottom: 20px; }
        .card { background: white; border-radius: 10px; padding: 24px; box-shadow: 0 2px 8px rgba(0,0,0,0.06); }
        .card.full { grid-column: 1 / -1; }
        .card h2 { font-size: 18px; color: #1a5276; margin-bottom: 16px; padding-bottom: 8px; border-bottom: 2px solid #eee; }
        .card h3 { font-size: 15px; color: #555; margin: 12px 0 8px; }
        table { width: 100%; border-collapse: collapse; font-size: 14px; }
        th { background: #f8f9fa; text-align: left; padding: 10px 12px; font-weight: 600; color: #555; }
        td { padding: 8px 12px; border-bottom: 1px solid #eee; }
        tr.sig td { background: #fff8e1; }
        .badge { padding: 2px 8px; border-radius: 4px; font-size: 12px; font-weight: 600; }
        .badge.tchem { background: #e8f5e9; color: #2e7d32; }
        .badge.tbio { background: #fff3e0; color: #e65100; }
        .badge.tclin { background: #e3f2fd; color: #1565c0; }
        .bar-container { display: flex; align-items: center; gap: 8px; }
        .bar { height: 20px; background: linear-gradient(90deg, #2980b9, #3498db); border-radius: 4px; min-width: 4px; }
        .kpi { display: flex; gap: 16px; flex-wrap: wrap; }
        .kpi-item { flex: 1; min-width: 120px; text-align: center; padding: 16px; background: #f8f9fa; border-radius: 8px; }
        .kpi-value { font-size: 28px; font-weight: 700; color: #1a5276; }
        .kpi-label { font-size: 12px; color: #888; margin-top: 4px; }
        .summary-box { background: #e8f5e9; border-left: 4px solid #4CAF50; padding: 16px; border-radius: 0 8px 8px 0; margin-top: 12px; }
        .warning-box { background: #fff3e0; border-left: 4px solid #FF9800; padding: 16px; border-radius: 0 8px 8px 0; margin-top: 12px; }
        .footer { text-align: center; color: #aaa; font-size: 12px; margin-top: 32px; padding: 16px; }
        .stage-tag { display: inline-block; background: #2980b9; color: white; padding: 2px 10px; border-radius: 12px; font-size: 12px; margin-bottom: 8px; }
        .stage-tag.s1 { background: #27ae60; }
        .stage-tag.s2 { background: #8e44ad; }
        .stage-tag.s3 { background: #e67e22; }
        .stage-tag.s4 { background: #c0392b; }
        .chart-img { max-width: 100%; height: auto; border-radius: 8px; margin: 12px 0; }
        .evidence-list { list-style: none; padding: 0; }
        .evidence-list li { padding: 8px 0; border-bottom: 1px solid #f0f0f0; }
        .evidence-list li::before { content: ""; display: inline-block; width: 8px; height: 8px; background: #2980b9; border-radius: 50%; margin-right: 8px; }
        @media (max-width: 768px) { .grid { grid-template-columns: 1fr; } }
        @media print { body { background: white; } .container { max-width: 100%; } }
    </style>
</head>
<body>
<div class="container">

    <div class="header">
        <h1>{{ title }}</h1>
        <div class="meta">
            生成时间: {{ generated_at }} &nbsp;|&nbsp;
            数据来源: {{ data_sources }} &nbsp;|&nbsp;
            Pipeline: drug-target-research-platform v1.0
        </div>
    </div>

    {% if kpis %}
    <div class="card full">
        <div class="kpi">
            {% for kpi in kpis %}
            <div class="kpi-item">
                <div class="kpi-value">{{ kpi.value }}</div>
                <div class="kpi-label">{{ kpi.label }}</div>
            </div>
            {% endfor %}
        </div>
    </div>
    {% endif %}

    {% for section in sections %}
    <div class="card {{ 'full' if section.get('full', false) else '' }}">
        {% if section.get('stage') %}
        <span class="stage-tag s{{ section.stage }}">Stage {{ section.stage }}</span>
        {% endif %}
        <h2>{{ section.title }}</h2>

        {% if section.get('chart') %}
        <img class="chart-img" src="{{ section.chart }}" alt="{{ section.title }}">
        {% endif %}

        {% if section.get('table') %}
        <table>
            <tr>
                {% for col in section.table.columns %}
                <th>{{ col }}</th>
                {% endfor %}
            </tr>
            {% for row in section.table.rows %}
            <tr class="{{ 'sig' if row.get('significant') else '' }}">
                {% for col in section.table.columns %}
                <td>{{ row.get(col, '') }}</td>
                {% endfor %}
            </tr>
            {% endfor %}
        </table>
        {% endif %}

        {% if section.get('summary') %}
        <div class="summary-box">{{ section.summary }}</div>
        {% endif %}

        {% if section.get('warning') %}
        <div class="warning-box">{{ section.warning }}</div>
        {% endif %}

        {% if section.get('content') %}
        {{ section.content }}
        {% endif %}
    </div>
    {% endfor %}

    {% if recommendations %}
    <div class="card full">
        <h2>综合评估与建议</h2>
        <div class="grid" style="grid-template-columns:1fr 1fr">
            <div>
                <h3>优势</h3>
                <ul class="evidence-list">
                    {% for item in recommendations.strengths %}
                    <li>{{ item }}</li>
                    {% endfor %}
                </ul>
            </div>
            <div>
                <h3>风险</h3>
                <ul class="evidence-list">
                    {% for item in recommendations.risks %}
                    <li>{{ item }}</li>
                    {% endfor %}
                </ul>
            </div>
        </div>
        {% if recommendations.get('conclusion') %}
        <div class="summary-box" style="margin-top:16px">
            <strong>结论:</strong> {{ recommendations.conclusion }}
        </div>
        {% endif %}
    </div>
    {% endif %}

    {% if provenance %}
    <div class="card full">
        <h2>数据溯源</h2>
        <table>
            <tr><th>数据源</th><th>版本/日期</th><th>访问方式</th><th>用途</th></tr>
            {% for p in provenance %}
            <tr>
                <td>{{ p.source }}</td>
                <td>{{ p.version }}</td>
                <td>{{ p.access }}</td>
                <td>{{ p.purpose }}</td>
            </tr>
            {% endfor %}
        </table>
    </div>
    {% endif %}

    <div class="footer">
        药物靶点智能研究平台 — 自动生成报告 &nbsp;|&nbsp;
        数据来源: GTEx, TCGA/GDC, GEO, OpenTargets, PubMed, ClinicalTrials.gov
    </div>
</div>
</body>
</html>"""


class ReportGenerator:
    """生成 HTML/PDF/DOCX 研究报告"""

    def __init__(self, output_dir: str):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.env = Environment(loader=BaseLoader())
        self.env.globals["now"] = datetime.now().strftime("%Y-%m-%d %H:%M")

    def _build_stage1_context(self, stage_data: dict, gene_symbol: str) -> list[dict]:
        """从 Stage1 数据构建报告上下文"""
        sections = []

        # 组织表达
        tissues = stage_data.get("tissue_expression", {}).get("gtex", {}).get("tissues", {})
        if tissues:
            sorted_t = sorted(tissues.items(), key=lambda x: x[1], reverse=True)[:10]
            sections.append({
                "stage": 1,
                "title": "正常组织表达 (GTEx)",
                "table": {
                    "columns": ["rank", "tissue", "tpm"],
                    "rows": [
                        {"rank": i + 1, "tissue": name, "tpm": f"{val:.1f}"}
                        for i, (name, val) in enumerate(sorted_t)
                    ]
                }
            })

        # 生存分析
        survival = stage_data.get("survival", {})
        if survival:
            rows = []
            for ct, data in survival.items():
                if isinstance(data, dict):
                    rows.append({
                        "cancer_type": ct,
                        "median_high": data.get("median_high", ""),
                        "median_low": data.get("median_low", ""),
                        "hr": data.get("hr", ""),
                        "p_value": data.get("p_value", ""),
                        "significant": data.get("p_value", 1) < 0.05,
                    })
            if rows:
                sections.append({
                    "stage": 1,
                    "title": "生存预后分析",
                    "table": {
                        "columns": ["cancer_type", "median_high", "median_low", "hr", "p_value"],
                        "rows": rows,
                    }
                })

        # 免疫相关性
        immune = stage_data.get("immune", {})
        for cancer_type, corr_data in immune.items():
            if isinstance(corr_data, dict) and "correlations" in corr_data:
                rows = []
                for cell, info in corr_data["correlations"].items():
                    pval = info.get("p_value", 1)
                    rows.append({
                        "immune_cell": cell.replace("_", " "),
                        "correlation": f"{info.get('correlation', 0):+.3f}",
                        "p_value": f"{pval:.2e}",
                        "significant": pval < 0.05,
                    })
                if rows:
                    sections.append({
                        "stage": 1,
                        "title": f"免疫微环境相关性 ({cancer_type})",
                        "table": {
                            "columns": ["immune_cell", "correlation", "p_value"],
                            "rows": rows,
                        }
                    })

        return sections

    def _build_stage2_context(self, stage_data: dict) -> list[dict]:
        """从 Stage2 数据构建报告上下文"""
        sections = []

        # 候选靶点排名
        targets = stage_data.get("top_targets", [])
        if targets:
            rows = []
            for i, t in enumerate(targets):
                rows.append({
                    "rank": i + 1,
                    "gene": t.get("gene", ""),
                    "composite_score": f"{t.get('composite_score', 0):.3f}",
                    "de_score": f"{t.get('de_score', 0):.3f}",
                    "druggability_score": f"{t.get('druggability_score', 0):.3f}",
                    "druggability_class": t.get("druggability_class", ""),
                })
            sections.append({
                "stage": 2,
                "title": "候选靶点综合排名",
                "full": True,
                "table": {
                    "columns": ["rank", "gene", "composite_score", "de_score", "druggability_score", "druggability_class"],
                    "rows": rows,
                }
            })

        return sections

    def _build_stage3_context(self, stage_data: dict) -> list[dict]:
        """从 Stage3 数据构建报告上下文"""
        sections = []
        developability = stage_data.get("developability", {})

        if developability:
            sections.append({
                "stage": 3,
                "title": "可开发性评估",
                "content": f"""
                <div class="kpi">
                    <div class="kpi-item">
                        <div class="kpi-value">{developability.get('score', 'N/A')}</div>
                        <div class="kpi-label">可开发性评分</div>
                    </div>
                    <div class="kpi-item">
                        <div class="kpi-value">{developability.get('level', 'N/A')}</div>
                        <div class="kpi-label">等级</div>
                    </div>
                </div>
                <p style="margin-top:12px">{developability.get('interpretation', '')}</p>
                """
            })

        recommendations = stage_data.get("recommendations", [])
        if recommendations:
            sections.append({
                "stage": 3,
                "title": "策略建议",
                "content": "<ul>" + "".join(f"<li>{r}</li>" for r in recommendations) + "</ul>"
            })

        return sections

    def _build_stage4_context(self, stage_data: dict) -> list[dict]:
        """从 Stage4 数据构建报告上下文"""
        sections = []
        milestones = stage_data.get("milestones", {})

        if milestones:
            sections.append({
                "stage": 4,
                "title": "开发里程碑",
                "content": f"""
                <div class="kpi">
                    <div class="kpi-item">
                        <div class="kpi-value">{milestones.get('total_duration', 'N/A')}</div>
                        <div class="kpi-label">总时长</div>
                    </div>
                    <div class="kpi-item">
                        <div class="kpi-value">{milestones.get('total_cost', 'N/A')}</div>
                        <div class="kpi-label">总成本</div>
                    </div>
                    <div class="kpi-item">
                        <div class="kpi-value">{milestones.get('success_rate', 'N/A')}</div>
                        <div class="kpi-label">综合成功率</div>
                    </div>
                </div>
                """
            })

        return sections

    def generate_html(
        self,
        pipeline_results: dict,
        gene_symbol: str,
        disease: str,
        title: str = None,
        provenance: list[dict] = None,
        chart_files: dict[str, str] = None,
    ) -> str:
        """从 pipeline 结果生成 HTML 报告"""
        sections = []

        stages = pipeline_results.get("stages", pipeline_results)

        # Stage1
        s1 = stages.get("stage1", stages.get("tissue_expression", {}))
        if s1:
            sections.extend(self._build_stage1_context(s1.get("summary", s1), gene_symbol))

        # Stage2
        s2 = stages.get("stage2", {})
        if s2:
            sections.extend(self._build_stage2_context(s2.get("summary", s2)))

        # Stage3
        s3 = stages.get("stage3", {})
        if s3:
            sections.extend(self._build_stage3_context(s3.get("summary", s3)))

        # Stage4
        s4 = stages.get("stage4", {})
        if s4:
            sections.extend(self._build_stage4_context(s4.get("summary", s4)))

        # KPI 概览
        kpis = []
        if gene_symbol:
            kpis.append({"value": gene_symbol, "label": "靶点"})
        if disease:
            kpis.append({"value": disease, "label": "疾病"})

        # 默认溯源
        if provenance is None:
            provenance = [
                {"source": "GTEx", "version": "v8", "access": "REST API", "purpose": "正常组织表达"},
                {"source": "TCGA/GDC", "version": "2024", "access": "GDC API", "purpose": "肿瘤表达与生存"},
                {"source": "GEO", "version": "2024", "access": "NCBI E-utilities", "purpose": "差异表达"},
                {"source": "OpenTargets", "version": "2024.02", "access": "GraphQL API", "purpose": "靶点-疾病关联"},
            ]

        template = self.env.from_string(HTML_TEMPLATE)
        html = template.render(
            title=title or f"{gene_symbol} — {disease} 靶点研究报告",
            generated_at=datetime.now().strftime("%Y-%m-%d %H:%M"),
            data_sources="GTEx, TCGA/GDC, GEO, OpenTargets",
            kpis=kpis,
            sections=sections,
            provenance=provenance,
        )

        path = self.output_dir / f"report_{gene_symbol}.html"
        path.write_text(html, encoding="utf-8")
        return str(path)

    def generate_pdf(self, html_path: str, filename: str = None) -> str:
        """HTML → PDF (WeasyPrint)"""
        try:
            from weasyprint import HTML
        except ImportError:
            raise RuntimeError("weasyprint not installed: pip install weasyprint")

        html_file = Path(html_path)
        if not filename:
            filename = html_file.stem + ".pdf"
        pdf_path = self.output_dir / filename

        HTML(filename=str(html_file)).write_pdf(str(pdf_path))
        return str(pdf_path)

    def generate_docx(
        self,
        pipeline_results: dict,
        gene_symbol: str,
        disease: str,
        title: str = None,
    ) -> str:
        """生成 DOCX 报告"""
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            raise RuntimeError("python-docx not installed: pip install python-docx")

        doc = Document()

        # 标题
        doc_title = doc.add_heading(title or f"{gene_symbol} — {disease} 靶点研究报告", level=0)
        doc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # 元信息
        meta = doc.add_paragraph()
        meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta_run = meta.add_run(
            f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n"
            f"数据来源: GTEx, TCGA/GDC, GEO, OpenTargets\n"
            f"Pipeline: drug-target-research-platform v1.0"
        )
        meta_run.font.size = Pt(9)
        meta_run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        doc.add_paragraph()  # spacer

        stages = pipeline_results.get("stages", pipeline_results)

        # Stage 1
        doc.add_heading("Stage 1: 表达与预后分析", level=1)
        s1 = stages.get("stage1", {}).get("summary", {})
        tissues = s1.get("tissue_expression", {}).get("gtex", {}).get("tissues", {})
        if tissues:
            doc.add_heading("正常组织表达 (GTEx)", level=2)
            sorted_t = sorted(tissues.items(), key=lambda x: x[1], reverse=True)[:10]
            table = doc.add_table(rows=1, cols=3, style="Light Grid Accent 1")
            table.rows[0].cells[0].text = "排名"
            table.rows[0].cells[1].text = "组织"
            table.rows[0].cells[2].text = "TPM"
            for i, (name, val) in enumerate(sorted_t):
                row = table.add_row()
                row.cells[0].text = str(i + 1)
                row.cells[1].text = name
                row.cells[2].text = f"{val:.1f}"

        # Stage 2
        doc.add_heading("Stage 2: 靶点筛选与排名", level=1)
        s2 = stages.get("stage2", {}).get("summary", {})
        targets = s2.get("top_targets", [])
        if targets:
            table = doc.add_table(rows=1, cols=5, style="Light Grid Accent 1")
            headers = ["排名", "靶点", "综合评分", "DE评分", "成药性"]
            for i, h in enumerate(headers):
                table.rows[0].cells[i].text = h
            for i, t in enumerate(targets):
                row = table.add_row()
                row.cells[0].text = str(i + 1)
                row.cells[1].text = t.get("gene", "")
                row.cells[2].text = f"{t.get('composite_score', 0):.3f}"
                row.cells[3].text = f"{t.get('de_score', 0):.3f}"
                row.cells[4].text = f"{t.get('druggability_score', 0):.3f}"

        # Stage 3
        doc.add_heading("Stage 3: 深度研究", level=1)
        s3 = stages.get("stage3", {}).get("summary", {})
        dev = s3.get("developability", {})
        if dev:
            doc.add_paragraph(f"可开发性评分: {dev.get('score', 'N/A')} ({dev.get('level', 'N/A')})")
            doc.add_paragraph(dev.get("interpretation", ""))

        # Stage 4
        doc.add_heading("Stage 4: 商业计划", level=1)
        s4 = stages.get("stage4", {}).get("summary", {})
        milestones = s4.get("milestones", {})
        if milestones:
            doc.add_paragraph(
                f"总时长: {milestones.get('total_duration', 'N/A')}\n"
                f"总成本: {milestones.get('total_cost', 'N/A')}\n"
                f"综合成功率: {milestones.get('success_rate', 'N/A')}"
            )

        # 溯源
        doc.add_heading("数据溯源", level=1)
        provenance = [
            ("GTEx", "v8", "REST API", "正常组织表达"),
            ("TCGA/GDC", "2024", "GDC API", "肿瘤表达与生存"),
            ("GEO", "2024", "NCBI E-utilities", "差异表达"),
            ("OpenTargets", "2024.02", "GraphQL API", "靶点-疾病关联"),
        ]
        table = doc.add_table(rows=1, cols=4, style="Light Grid Accent 1")
        for i, h in enumerate(["数据源", "版本", "访问方式", "用途"]):
            table.rows[0].cells[i].text = h
        for source, version, access, purpose in provenance:
            row = table.add_row()
            row.cells[0].text = source
            row.cells[1].text = version
            row.cells[2].text = access
            row.cells[3].text = purpose

        path = self.output_dir / f"report_{gene_symbol}.docx"
        doc.save(str(path))
        return str(path)
